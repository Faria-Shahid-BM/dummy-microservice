"""Tests for the structural Document Diff engine.

Fixtures are built with ``python-docx`` at run time rather than committed as
binaries, so a reader can see exactly what each case contains and change it
without a round trip through Word.
"""
from __future__ import annotations

import re

import pytest
from bs4 import BeautifulSoup
from docx import Document

from engines import document_diff_html as dd


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def para(*texts: str) -> str:
    """Minimal HTML the way mammoth emits it, for token-level cases."""
    return "".join(f"<p>{t}</p>" for t in texts)


def marked_text(html: str, kind: str = "delete") -> str:
    """The visible text that actually renders struck-through / inserted.

    Read off the rendered markup rather than the ``changes`` list, because the
    two can disagree: the bug :func:`dd._wrap_runs` fixes was a change being
    reported correctly while the redline showed it unmarked.
    """
    soup = BeautifulSoup(html, "html.parser")
    inline, block = f"seg-{kind}", f"seg-{kind}-block"
    out = []
    for el in soup.find_all(True):
        classes = el.get("class") or []
        if inline in classes or block in classes:
            if any(block in (a.get("class") or []) for a in el.parents):
                continue  # already counted via the ancestor row
            out.append(re.sub(r"\s+", " ", el.get_text(" ", strip=True)))
    return " ".join(t for t in out if t)


def build_docx(path, rows=None, table=True, tail="This Schedule forms part of the Contract."):
    doc = Document()
    doc.add_paragraph("Schedule of Secured Assets")
    doc.add_paragraph("The Mortgagor charges the following assets to the Bank.")
    if table and rows:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.add_paragraph(tail)
    doc.save(str(path))
    return path


ROWS = [
    ["Asset", "Description", "Value (AED)"],
    ["1", "Villa, Plot 221, Jumeirah", "25,000,000"],
    ["2", "Warehouse, Al Quoz", "7,500,000"],
    ["3", "Office Unit, Business Bay", "3,250,000"],
]


# --------------------------------------------------------------------------
# tokenizer
# --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "html",
    [
        "<p>Plain sentence.</p>",
        "<p>Amount: 25,000,000 (AED); rate 2.5% p.a.</p>",
        "<table><tr><td><p>a</p></td><td><p>b</p></td></tr></table>",
        "<p>الضمان الكامل والانفراد.</p>",
        "<p>Tabs\tand\n newlines  collapse?</p>",
        "<p>Hyphen-ated, it's o'clock, and/or 1,234.56</p>",
    ],
)
def test_tokenizer_is_lossless(html):
    """The redline is rebuilt by joining tokens, so a lossy regex corrupts it."""
    assert "".join(dd._HTML_TOKEN_RE.findall(html)) == html


def test_punctuation_splits_off_words():
    assert dd._HTML_TOKEN_RE.findall("<p>Guarantee.</p>") == ["<p>", "Guarantee", ".", "</p>"]


@pytest.mark.parametrize("value", ["25,000,000", "1,234.56", "2.5", "3.250", "0.05"])
def test_numbers_stay_one_token(value):
    """Splitting punctuation must not shatter amounts into digit runs."""
    assert dd._HTML_TOKEN_RE.findall(f"<p>{value}</p>") == ["<p>", value, "</p>"]


def test_internal_word_punctuation_stays_joined():
    assert dd._HTML_TOKEN_RE.findall("<p>and/or</p>") == ["<p>", "and/or", "</p>"]
    assert dd._HTML_TOKEN_RE.findall("<p>Hyphen-ated</p>") == ["<p>", "Hyphen-ated", "</p>"]


# --------------------------------------------------------------------------
# the punctuation-boundary bug: truncated clause is a deletion, not a
# replacement, and the surviving word is not dragged into the change
# --------------------------------------------------------------------------

def test_truncated_sentence_is_a_deletion_not_a_replacement():
    original = para(
        "The provisions of Article 1092 are not applicable to this Guarantee "
        "and the Bank shall not be obliged to make any demand within the six "
        "month period mentioned in that Article."
    )
    returned = para("The provisions of Article 1092 are not applicable to this Guarantee.")

    result = dd.compare_documents_html(original, returned)

    assert result["summary"] == {
        "insertions": 0,
        "deletions": 1,
        "replacements": 0,
        "changes": 1,
    }
    change = result["changes"][0]
    assert change["type"] == "deletion"
    assert change["after"] == ""
    # the word that survived in both documents must not appear in the change
    assert not change["before"].startswith("Guarantee")
    assert change["before"] == (
        "and the Bank shall not be obliged to make any demand within the six "
        "month period mentioned in that Article"
    )


def test_trailing_clause_removal_leaves_no_stranded_punctuation():
    original = para("The maximum amount shall be 25,000,000 (AED) together with any additional amount.")
    returned = para("The maximum amount shall be 25,000,000 (AED).")

    result = dd.compare_documents_html(original, returned)

    assert [c["type"] for c in result["changes"]] == ["deletion"]
    # a bare "." as the inserted side was the old symptom
    assert all(c["after"] == "" for c in result["changes"])
    assert result["changes"][0]["before"] == "together with any additional amount"


def test_amount_change_is_one_replacement_with_whole_numbers():
    original = para("The maximum amount guaranteed shall be 25,000,000 (AED).")
    returned = para("The maximum amount guaranteed shall be 30,500,000 (AED).")

    result = dd.compare_documents_html(original, returned)

    assert result["summary"]["changes"] == 1
    change = result["changes"][0]
    assert change["type"] == "replacement"
    assert change["before"] == "25,000,000"
    assert change["after"] == "30,500,000"


def test_negation_change_is_detected():
    """The stopword-stripping alternative would erase exactly this."""
    original = para("The Bank shall not be obliged to make any demand.")
    returned = para("The Bank shall be obliged to make any demand.")

    result = dd.compare_documents_html(original, returned)

    assert result["summary"]["changes"] == 1
    assert result["changes"][0]["before"] == "not"


def test_identical_documents():
    html = para("Clause one.", "Clause two, with 25,000,000 (AED).")
    result = dd.compare_documents_html(html, html)
    assert result["identical"] is True
    assert result["similarity"] == pytest.approx(1.0)
    assert result["changes"] == []


def test_large_removal_is_flagged_as_possible_missing_section():
    body = " ".join(f"word{i}" for i in range(dd.LARGE_SECTION_WORD_THRESHOLD + 5))
    original = para("Keep this.", body, "Keep that.")
    returned = para("Keep this.", "Keep that.")

    result = dd.compare_documents_html(original, returned)

    assert any(c["possibleMissingSection"] for c in result["changes"])


# --------------------------------------------------------------------------
# structural scenarios: a missing row and a missing table
# --------------------------------------------------------------------------

def test_missing_row_is_fully_marked_in_the_redline(tmp_path):
    """Regression: only the first cell of a removed row used to be struck."""
    original = dd.docx_to_html(build_docx(tmp_path / "full.docx", ROWS)).html
    returned = dd.docx_to_html(
        build_docx(tmp_path / "missing_row.docx", [ROWS[0], ROWS[1], ROWS[3]])
    ).html

    result = dd.compare_documents_html(original, returned)

    assert result["summary"]["deletions"] == 1
    reported = result["changes"][0]["before"]
    assert reported == "2 Warehouse, Al Quoz 7,500,000"
    # what the reviewer actually sees must match what the changes list claims
    assert marked_text(result["html"]) == reported


def test_missing_row_gets_the_block_class_the_frontend_styles(tmp_path):
    original = dd.docx_to_html(build_docx(tmp_path / "full.docx", ROWS)).html
    returned = dd.docx_to_html(
        build_docx(tmp_path / "missing_row.docx", [ROWS[0], ROWS[1], ROWS[3]])
    ).html

    result = dd.compare_documents_html(original, returned)
    soup = BeautifulSoup(result["html"], "html.parser")

    rows = soup.select("tr.seg-delete-block")
    assert len(rows) == 1
    assert "Warehouse" in rows[0].get_text(" ")
    # click-to-jump anchor rides along on the row
    assert any(c.startswith("doc-change-") for c in rows[0].get("class"))
    # the surviving rows are untouched
    assert len(soup.find_all("tr")) == len(ROWS)


def test_missing_table_is_reported_and_fully_marked(tmp_path):
    original = dd.docx_to_html(build_docx(tmp_path / "full.docx", ROWS)).html
    returned = dd.docx_to_html(build_docx(tmp_path / "no_table.docx", ROWS, table=False)).html

    result = dd.compare_documents_html(original, returned)

    assert result["summary"]["deletions"] == 1
    change = result["changes"][0]
    for cell in ("Villa, Plot 221, Jumeirah", "25,000,000", "Office Unit, Business Bay"):
        assert cell in change["before"]
    assert marked_text(result["html"]) == change["before"]
    assert change["possibleMissingSection"] is False  # 21 words, below the threshold


def test_missing_table_survives_the_word_threshold_when_large(tmp_path):
    """A table big enough to clear the threshold is flagged as a lost section."""
    big = [ROWS[0]] + [[str(i), f"Asset number {i} at some location", f"{i},000,000"] for i in range(1, 12)]
    original = dd.docx_to_html(build_docx(tmp_path / "big.docx", big)).html
    returned = dd.docx_to_html(build_docx(tmp_path / "big_none.docx", big, table=False)).html

    result = dd.compare_documents_html(original, returned)

    assert any(c["possibleMissingSection"] for c in result["changes"])


def test_partial_row_edit_is_not_marked_as_a_whole_row(tmp_path):
    """Only a changed value in one cell — the row must stay un-shaded."""
    edited = [row[:] for row in ROWS]
    edited[2][2] = "9,900,000"
    original = dd.docx_to_html(build_docx(tmp_path / "full.docx", ROWS)).html
    returned = dd.docx_to_html(build_docx(tmp_path / "edited.docx", edited)).html

    result = dd.compare_documents_html(original, returned)
    soup = BeautifulSoup(result["html"], "html.parser")

    assert soup.select("tr.seg-delete-block") == []
    assert result["changes"][0]["before"] == "7,500,000"
    assert result["changes"][0]["after"] == "9,900,000"


def test_redline_html_is_well_formed(tmp_path):
    """Re-parsing the emitted redline must not move or drop any table row."""
    original = dd.docx_to_html(build_docx(tmp_path / "full.docx", ROWS)).html
    returned = dd.docx_to_html(
        build_docx(tmp_path / "missing_row.docx", [ROWS[0], ROWS[3]])
    ).html

    result = dd.compare_documents_html(original, returned)
    soup = BeautifulSoup(result["html"], "html.parser")

    assert len(soup.find_all("table")) == 1
    # every row still sits inside the table; none foster-parented out
    assert all(row.find_parent("table") is not None for row in soup.find_all("tr"))


# --------------------------------------------------------------------------
# images: dropped before the diff, but counted and reported
# --------------------------------------------------------------------------

def make_png(path, seed):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (300, 100), "white")
    draw = ImageDraw.Draw(img)
    for x in range(0, 300, 7):
        draw.line([(x, 50 + (x * seed) % 30), (x + 7, 50 + ((x + 13) * seed) % 30)], fill="navy", width=3)
    img.save(path)
    return path


def build_signed(path, sig=None, amount="25,000,000"):
    """The real shape of a Document Reviewer pair: the generated original has
    no signature, the returned copy does."""
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph("CORPORATE GUARANTEE")
    doc.add_paragraph(f"The maximum amount guaranteed shall be {amount} (AED).")
    doc.add_paragraph("Signed by the authorised signatory:")
    if sig is not None:
        doc.add_picture(str(sig), width=Inches(2.0))
    doc.save(str(path))
    return path


def test_images_are_dropped_and_counted(tmp_path):
    sig = make_png(tmp_path / "sig.png", 3)
    converted = dd.docx_to_html(build_signed(tmp_path / "signed.docx", sig))

    assert converted.image_count == 1
    assert "<img" not in converted.html
    assert "base64" not in converted.html  # never materialised, not just stripped
    assert "CORPORATE GUARANTEE" in converted.html


def test_generated_original_reports_no_images(tmp_path):
    assert dd.docx_to_html(build_signed(tmp_path / "gen.docx")).image_count == 0


def test_signing_an_unchanged_document_is_still_identical(tmp_path):
    """The returned copy always gains a signature the original cannot have;
    that must not read as a change."""
    sig = make_png(tmp_path / "sig.png", 3)
    original = dd.docx_to_html(build_signed(tmp_path / "gen.docx")).html
    returned = dd.docx_to_html(build_signed(tmp_path / "signed.docx", sig)).html

    result = dd.compare_documents_html(original, returned)

    assert result["identical"] is True
    assert result["summary"]["changes"] == 0


def test_tamper_is_still_caught_under_a_signature(tmp_path):
    sig = make_png(tmp_path / "sig.png", 3)
    original = dd.docx_to_html(build_signed(tmp_path / "gen.docx")).html
    returned = dd.docx_to_html(
        build_signed(tmp_path / "signed.docx", sig, amount="99,000,000")
    ).html

    result = dd.compare_documents_html(original, returned)

    assert result["summary"]["changes"] == 1
    assert result["changes"][0]["before"] == "25,000,000"
    assert result["changes"][0]["after"] == "99,000,000"


def test_dropping_images_keeps_the_redline_small(tmp_path):
    """A 2 MB scan used to inline as ~2.6 MB of base64 into the stored html."""
    from PIL import Image
    import random

    random.seed(0)
    big = Image.new("RGB", (1200, 400))
    big.putdata([(random.randrange(256),) * 3 for _ in range(1200 * 400)])
    big.save(tmp_path / "scan.png")

    original = dd.docx_to_html(build_signed(tmp_path / "gen.docx")).html
    returned = dd.docx_to_html(
        build_signed(tmp_path / "signed.docx", tmp_path / "scan.png", amount="99,000,000")
    ).html

    result = dd.compare_documents_html(original, returned)

    assert (tmp_path / "scan.png").stat().st_size > 500_000
    assert len(result["html"]) < 10_000


def test_scanned_document_with_no_text_layer_yields_nothing(tmp_path):
    """The service turns this into a readable 'No text extracted' error."""
    from docx.shared import Inches

    doc = Document()
    doc.add_picture(str(make_png(tmp_path / "scan.png", 5)), width=Inches(4))
    doc.save(str(tmp_path / "scanonly.docx"))

    converted = dd.docx_to_html(tmp_path / "scanonly.docx")

    assert converted.image_count == 1
    assert dd.html_to_text(converted.html) == ""
